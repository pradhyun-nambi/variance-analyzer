from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

COLUMNS = ["month", "line_item", "type", "budget", "actual",
           "variance", "variance_pct", "status"]


def build_report(df, bullets, title="Budget vs. Actual Variance Report"):
    """Return a one-page Excel report as bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Variance report"
    bold = Font(bold=True)

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "SYNTHETIC data for a hypothetical business. Not real financials."

    ws["A4"] = "Key points"
    ws["A4"].font = bold
    row = 5
    for b in bullets:
        ws.cell(row=row, column=1, value="- " + b)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Flagged items").font = bold
    row += 1
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    row += 1

    for _, r in df[df["flagged"]].iterrows():
        for c, name in enumerate(COLUMNS, start=1):
            v = r[name]
            if pd.isna(v):
                v = None
            elif name in ("budget", "actual", "variance", "variance_pct"):
                v = float(v)
            cell = ws.cell(row=row, column=c, value=v)
            if name in ("budget", "actual", "variance"):
                cell.number_format = '$#,##0;[Red]-$#,##0'
            elif name == "variance_pct":
                cell.number_format = '0.0"%"'
        row += 1

    for col, width in zip("ABCDEFGH", [10, 24, 10, 10, 10, 10, 13, 12]):
        ws.column_dimensions[col].width = width

    # Print on one landscape page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()