from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from src.variance import add_variance, flag_material
from src.export import build_report


def build(rows):
    df = pd.DataFrame(rows, columns=["month", "line_item", "type", "budget", "actual"])
    return flag_material(add_variance(df))


def test_report_has_title_and_synthetic_label():
    df = build([["2026-05", "Equipment rental", "expense", 350, 1100]])
    ws = load_workbook(BytesIO(build_report(df, ["test bullet"]))).active
    assert ws["A1"].value == "Budget vs. Actual Variance Report"
    assert "SYNTHETIC" in ws["A2"].value


def test_report_contains_flagged_row():
    df = build([["2026-05", "Equipment rental", "expense", 350, 1100]])
    ws = load_workbook(BytesIO(build_report(df, []))).active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "Equipment rental" in values
    assert 750.0 in values